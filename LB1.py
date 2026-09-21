import openpyxl
import time

PATH = '/Users/arsenijsaverskij/Documents/3 семестр/Algoritmu/BD_ALG.xlsx'
wb = openpyxl.load_workbook(PATH, data_only=True)
ws = wb.active
FIRST_ROW = 2
LAST_ROW = ws.max_row


def to_key(v):
    if v is None:
        return (2, "")
    if isinstance(v, bool):
        return (1, str(v))
    if isinstance(v, (int, float)):
        return (0, float(v))
    s = str(v)
    for ch in ("\xa0", "\u200b", "\ufeff", "\t", "\n", "\r"):
        s = s.replace(ch, "")
    s = s.strip()
    if s == "":
        return (2, "")
    try:
        return (0, float(s))
    except ValueError:
        return (1, s)


# --- читаем данные один раз ---
ARTICLES_KEY = []
VALUES = []
for r in range(FIRST_ROW, LAST_ROW + 1):
    ARTICLES_KEY.append(to_key(ws.cell(row=r, column=1).value))
    VALUES.append(ws.cell(row=r, column=2).value)

N = len(ARTICLES_KEY)

def code(input_field):
    target = to_key(input_field)

    # ---------- Линейный ----------
    start = time.perf_counter()
    found_index = None
    for i in range(N):
        if ARTICLES_KEY[i] == target:
            found_index = i
            break
    time_linear = time.perf_counter() - start

    print(VALUES[found_index] if found_index is not None else "Артикул не найден")
    print("|Линейный алгоритм|")
    print("Время выполнения кода:", time_linear)

    # ---------- Бинарный ----------
    left, right = 0, N - 1
    start = time.perf_counter()
    found_index = None
    while left <= right:
        mid = (left + right) // 2
        mid_key = ARTICLES_KEY[mid]
        if mid_key == target:
            found_index = mid
            break
        elif target < mid_key:
            right = mid - 1
        else:
            left = mid + 1
    time_binary = time.perf_counter() - start

    print(VALUES[found_index] if found_index is not None else "Артикул не найден")
    print("|Бинарный код|")
    print("Время выполнения кода:", time_binary)

    # ---------- Экспоненциальный ----------
    start = time.perf_counter()
    found_index = None
    if N > 0:
        bound = 1
        while bound < N and ARTICLES_KEY[bound] < target:
            bound *= 2
        left2 = bound // 2
        right2 = min(bound, N - 1)
        while left2 <= right2:
            mid = (left2 + right2) // 2
            mid_key = ARTICLES_KEY[mid]
            if mid_key == target:
                found_index = mid
                break
            elif mid_key < target:
                left2 = mid + 1
            else:
                right2 = mid - 1
    time_exp = time.perf_counter() - start

    print(VALUES[found_index] if found_index is not None else "Артикул не найден")
    print("|Экспоненциальный способ|")
    print("Время выполнения кода:", time_exp)

    # ---------- Прыжковый ----------
    start = time.perf_counter()
    found_index = None
    if N > 0:
        step = max(1, int(N ** 0.5))
        prev, curr = 0, step
        while curr < N and ARTICLES_KEY[curr] < target:
            prev = curr
            curr += step
        end_block = min(curr, N)
        for i in range(prev, end_block):
            if ARTICLES_KEY[i] == target:
                found_index = i
                break
    time_jump = time.perf_counter() - start

    print(VALUES[found_index] if found_index is not None else "Артикул не найден")
    print("|Прыжковый метод|")
    print("Время выполнения кода:", time_jump)


    times = {
        "линейный": time_linear,
        "бинарный": time_binary,
        "экспоненциальный": time_exp,
        "прыжковый": time_jump,
    }
    best_time = min(times.values())
    fastest_list = [k for k, v in times.items() if abs(v - best_time) < 1e-9]

    if len(fastest_list) == 1:
        print("Самый быстрый:", fastest_list[0])
    else:
        print("Самые быстрые:", ", ".join(fastest_list), "(равны)")


# ---------- Тесты ----------
code(1)
print("1------------------------------------")
code(2)
print("2------------------------------------")
code(194056)
print("3------------------------------------")
code(450000)
print("4------------------------------------")
code(890000)
print("5------------------------------------")
code(10000000000000000)
print("-------------------------------------")